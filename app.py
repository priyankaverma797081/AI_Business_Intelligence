from flask import Flask, render_template, jsonify, request, Response, send_file
import mysql.connector
import pandas as pd
import os
import io
from copy import copy

from openpyxl import Workbook

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak
)


app = Flask(__name__)


# =========================================================
# DATABASE CONFIGURATION
# =========================================================

from dotenv import load_dotenv

load_dotenv()

DB_CONFIG = {
    "host": os.getenv("MYSQLHOST", "localhost"),
    "port": int(os.getenv("MYSQLPORT", "3306")),
    "user": os.getenv("MYSQLUSER", "root"),
    "password": os.getenv("MYSQLPASSWORD"),
    "database": os.getenv(
        "MYSQLDATABASE",
        "ai_business_intelligence"
    )
}


# =========================================================
# DATABASE CONNECTION
# =========================================================

def get_db_connection():
    return mysql.connector.connect(**DB_CONFIG)


# =========================================================
# BUILD FILTERS
# =========================================================

def build_sales_filters():
    conditions = []
    params = []

    category = request.args.get("category", "").strip()
    region = request.args.get("region", "").strip()
    month = request.args.get("month", "").strip()

    if category and category != "All":
        conditions.append("Category = %s")
        params.append(category)

    if region and region != "All":
        conditions.append("Region = %s")
        params.append(region)

    if month and month != "All":
        conditions.append(
            "DATE_FORMAT(Order_Date, '%Y-%m') = %s"
        )
        params.append(month)

    if conditions:
        filters = "WHERE " + " AND ".join(conditions)
    else:
        filters = ""

    return filters, params


# =========================================================
# HOME
# =========================================================

@app.route("/")
def index():
    return render_template("index.html")


# =========================================================
# HEALTH CHECK
# =========================================================

@app.route("/api/health")
def health():

    try:
        conn = get_db_connection()
        conn.close()

        return jsonify({
            "status": "success",
            "message": "Database connection successful"
        })

    except Exception as e:
        return jsonify({
            "status": "error",
            "message": str(e)
        }), 500


# =========================================================
# KPI API
# =========================================================

@app.route("/api/kpis")
def get_kpis():

    try:
        filters, params = build_sales_filters()

        query = f"""
            SELECT
                COALESCE(SUM(Total_Sales), 0) AS total_sales,
                COUNT(*) AS total_orders,
                COUNT(DISTINCT Customer_ID) AS total_customers,
                COALESCE(AVG(Total_Sales), 0) AS average_order_value
            FROM sales
            {filters}
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        result = cursor.fetchone()

        cursor.close()
        conn.close()

        return jsonify({
            "total_sales": float(result["total_sales"]),
            "total_orders": int(result["total_orders"]),
            "total_customers": int(result["total_customers"]),
            "average_order_value": float(
                result["average_order_value"]
            )
        })

    except Exception as e:
        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# SALES API
# =========================================================

@app.route("/api/sales")
def get_sales():

    try:
        filters, params = build_sales_filters()

        query = f"""
            SELECT
                Order_ID,
                Order_Date,
                Customer_ID,
                Product,
                Category,
                Region,
                Quantity,
                Unit_Price,
                Discount_Pct,
                Total_Sales
            FROM sales
            {filters}
            ORDER BY Order_Date DESC
            LIMIT 100
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        for row in rows:

            if row["Order_Date"]:
                row["Order_Date"] = (
                    row["Order_Date"].strftime("%Y-%m-%d")
                )

            row["Unit_Price"] = float(
                row["Unit_Price"]
            )

            row["Discount_Pct"] = float(
                row["Discount_Pct"]
            )

            row["Total_Sales"] = float(
                row["Total_Sales"]
            )

        return jsonify(rows)

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# CATEGORY SALES
# =========================================================

@app.route("/api/category-sales")
def get_category_sales():

    try:
        filters, params = build_sales_filters()

        query = f"""
            SELECT
                Category,
                SUM(Total_Sales) AS total_sales
            FROM sales
            {filters}
            GROUP BY Category
            ORDER BY total_sales DESC
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        for row in rows:
            row["total_sales"] = float(
                row["total_sales"]
            )

        return jsonify(rows)

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# REGION SALES
# =========================================================

@app.route("/api/region-sales")
def get_region_sales():

    try:
        filters, params = build_sales_filters()

        query = f"""
            SELECT
                Region,
                SUM(Total_Sales) AS total_sales
            FROM sales
            {filters}
            GROUP BY Region
            ORDER BY total_sales DESC
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        for row in rows:
            row["total_sales"] = float(
                row["total_sales"]
            )

        return jsonify(rows)

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# MONTHLY SALES
# =========================================================

@app.route("/api/monthly-sales")
def get_monthly_sales():

    try:
        filters, params = build_sales_filters()

        query = f"""
            SELECT
                DATE_FORMAT(Order_Date, '%Y-%m') AS month,
                SUM(Total_Sales) AS total_sales
            FROM sales
            {filters}
            GROUP BY DATE_FORMAT(Order_Date, '%Y-%m')
            ORDER BY month
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        for row in rows:
            row["total_sales"] = float(
                row["total_sales"]
            )

        return jsonify(rows)

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# FORECAST API
# =========================================================

@app.route("/api/forecast")
def get_forecast():

    try:

        forecast_path = os.path.join(
            "data",
            "sales_forecast.csv"
        )

        if not os.path.exists(forecast_path):

            return jsonify({
                "error": "Forecast file not found"
            }), 404

        df = pd.read_csv(forecast_path)

        return jsonify(
            df.to_dict(orient="records")
        )

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# FORECAST INSIGHTS
# =========================================================

@app.route("/api/forecast-insights")
def forecast_insights():

    try:

        forecast_path = os.path.join(
            "data",
            "sales_forecast.csv"
        )

        if not os.path.exists(forecast_path):

            return jsonify({
                "error": "Forecast file not found"
            }), 404

        df = pd.read_csv(forecast_path)

        if df.empty:

            return jsonify({
                "error": "Forecast data is empty"
            }), 404

        possible_columns = [
            "Forecast_Sales",
            "Predicted_Sales",
            "Forecast",
            "Sales_Forecast"
        ]

        sales_column = None

        for column in possible_columns:

            if column in df.columns:
                sales_column = column
                break

        if sales_column is None:

            numeric_columns = df.select_dtypes(
                include="number"
            ).columns.tolist()

            if numeric_columns:
                sales_column = numeric_columns[-1]

        if sales_column is None:

            return jsonify({
                "error": "Forecast sales column not found"
            }), 500

        forecast_values = pd.to_numeric(
            df[sales_column],
            errors="coerce"
        ).dropna()

        if forecast_values.empty:

            return jsonify({
                "error": "No valid forecast values found"
            }), 500

        next_month = float(
            forecast_values.iloc[0]
        )

        average_forecast = float(
            forecast_values.mean()
        )

        highest_forecast = float(
            forecast_values.max()
        )

        highest_index = forecast_values.idxmax()

        highest_month = str(
            df.loc[highest_index].iloc[0]
        )

        if len(forecast_values) >= 2:

            first_value = float(
                forecast_values.iloc[0]
            )

            last_value = float(
                forecast_values.iloc[-1]
            )

            if first_value != 0:

                overall_trend = (
                    (last_value - first_value)
                    / first_value
                ) * 100

            else:

                overall_trend = 0

        else:

            overall_trend = 0

        return jsonify({

            "next_month_forecast": next_month,

            "average_forecast": average_forecast,

            "highest_forecast": highest_forecast,

            "highest_forecast_month": highest_month,

            "overall_trend": round(
                overall_trend,
                2
            )
        })

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# FILTER OPTIONS
# =========================================================

@app.route("/api/filter-options")
def filter_options():

    try:

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        # Categories
        cursor.execute("""
            SELECT DISTINCT Category
            FROM sales
            ORDER BY Category
        """)

        categories = [
            row["Category"]
            for row in cursor.fetchall()
        ]

        # Regions
        cursor.execute("""
            SELECT DISTINCT Region
            FROM sales
            ORDER BY Region
        """)

        regions = [
            row["Region"]
            for row in cursor.fetchall()
        ]

        # Months
        cursor.execute("""
            SELECT DISTINCT
                DATE_FORMAT(
                    Order_Date,
                    '%Y-%m'
                ) AS month
            FROM sales
            ORDER BY month
        """)

        months = [
            row["month"]
            for row in cursor.fetchall()
        ]

        cursor.close()
        conn.close()

        return jsonify({

            "categories": categories,

            "regions": regions,

            "months": months

        })

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# CSV EXPORT
# =========================================================

@app.route("/export/csv")
def export_csv():

    try:

        filters, params = build_sales_filters()

        query = f"""
            SELECT
                Order_ID,
                Order_Date,
                Customer_ID,
                Product,
                Category,
                Region,
                Quantity,
                Unit_Price,
                Discount_Pct,
                Total_Sales
            FROM sales
            {filters}
            ORDER BY Order_Date DESC
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        df = pd.DataFrame(rows)

        output = io.StringIO()

        df.to_csv(
            output,
            index=False
        )

        response = Response(
            output.getvalue(),
            mimetype="text/csv"
        )

        response.headers["Content-Disposition"] = (
            "attachment; filename=sales_report.csv"
        )

        return response

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# EXCEL EXPORT
# =========================================================

@app.route("/export/excel")
def export_excel():

    try:

        filters, params = build_sales_filters()

        query = f"""
            SELECT
                Order_ID,
                Order_Date,
                Customer_ID,
                Product,
                Category,
                Region,
                Quantity,
                Unit_Price,
                Discount_Pct,
                Total_Sales
            FROM sales
            {filters}
            ORDER BY Order_Date DESC
        """

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(query, params)
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Sales Report"

        headers = [
            "Order_ID",
            "Order_Date",
            "Customer_ID",
            "Product",
            "Category",
            "Region",
            "Quantity",
            "Unit_Price",
            "Discount_Pct",
            "Total_Sales"
        ]

        worksheet.append(headers)

        for row in rows:

            worksheet.append([
                row["Order_ID"],
                row["Order_Date"],
                row["Customer_ID"],
                row["Product"],
                row["Category"],
                row["Region"],
                row["Quantity"],
                row["Unit_Price"],
                row["Discount_Pct"],
                row["Total_Sales"]
            ])

        # Header formatting
        for cell in worksheet[1]:

            new_font = copy(cell.font)
            new_font.bold = True
            cell.font = new_font

        # Number formatting
        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row
        ):

            row[7].number_format = '#,##0.00'
            row[8].number_format = '0.00'
            row[9].number_format = '#,##0.00'

        # Auto column width
        for column in worksheet.columns:

            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = max_length + 2

        excel_file = io.BytesIO()

        workbook.save(excel_file)

        excel_file.seek(0)

        return send_file(

            excel_file,

            as_attachment=True,

            download_name="sales_report.xlsx",

            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# PDF BUSINESS REPORT
# =========================================================

@app.route("/export/pdf")
def export_pdf():

    try:

        # -------------------------------------------------
        # Current filters
        # -------------------------------------------------

        filters, params = build_sales_filters()

        category = request.args.get(
            "category",
            "All"
        )

        region = request.args.get(
            "region",
            "All"
        )

        month = request.args.get(
            "month",
            "All"
        )

        # -------------------------------------------------
        # KPI DATA
        # -------------------------------------------------

        kpi_query = f"""
            SELECT

                COALESCE(
                    SUM(Total_Sales),
                    0
                ) AS total_sales,

                COUNT(*) AS total_orders,

                COUNT(
                    DISTINCT Customer_ID
                ) AS total_customers,

                COALESCE(
                    AVG(Total_Sales),
                    0
                ) AS average_order_value

            FROM sales

            {filters}
        """

        conn = get_db_connection()

        cursor = conn.cursor(
            dictionary=True
        )

        cursor.execute(
            kpi_query,
            params
        )

        kpis = cursor.fetchone()

        # -------------------------------------------------
        # CATEGORY DATA
        # -------------------------------------------------

        category_query = f"""
            SELECT
                Category,
                SUM(Total_Sales) AS total_sales
            FROM sales
            {filters}
            GROUP BY Category
            ORDER BY total_sales DESC
        """

        cursor.execute(
            category_query,
            params
        )

        category_rows = cursor.fetchall()

        # -------------------------------------------------
        # REGION DATA
        # -------------------------------------------------

        region_query = f"""
            SELECT
                Region,
                SUM(Total_Sales) AS total_sales
            FROM sales
            {filters}
            GROUP BY Region
            ORDER BY total_sales DESC
        """

        cursor.execute(
            region_query,
            params
        )

        region_rows = cursor.fetchall()

        # -------------------------------------------------
        # MONTHLY DATA
        # -------------------------------------------------

        monthly_query = f"""
            SELECT

                DATE_FORMAT(
                    Order_Date,
                    '%Y-%m'
                ) AS month,

                SUM(Total_Sales)
                    AS total_sales

            FROM sales

            {filters}

            GROUP BY
                DATE_FORMAT(
                    Order_Date,
                    '%Y-%m'
                )

            ORDER BY month
        """

        cursor.execute(
            monthly_query,
            params
        )

        monthly_rows = cursor.fetchall()

        cursor.close()
        conn.close()

        # -------------------------------------------------
        # FORECAST DATA
        # -------------------------------------------------

        forecast_path = os.path.join(
            "data",
            "sales_forecast.csv"
        )

        forecast_df = pd.DataFrame()

        if os.path.exists(
            forecast_path
        ):

            forecast_df = pd.read_csv(
                forecast_path
            )

        # -------------------------------------------------
        # FIND FORECAST COLUMN
        # -------------------------------------------------

        forecast_values = []

        if not forecast_df.empty:

            possible_columns = [

                "Forecast_Sales",

                "Predicted_Sales",

                "Forecast",

                "Sales_Forecast"

            ]

            forecast_column = None

            for column in possible_columns:

                if column in forecast_df.columns:

                    forecast_column = column
                    break

            if forecast_column is None:

                numeric_columns = (
                    forecast_df
                    .select_dtypes(
                        include="number"
                    )
                    .columns
                    .tolist()
                )

                if numeric_columns:

                    forecast_column = (
                        numeric_columns[-1]
                    )

            if forecast_column:

                forecast_values = (

                    pd.to_numeric(

                        forecast_df[
                            forecast_column
                        ],

                        errors="coerce"

                    )
                    .dropna()
                    .tolist()
                )

        # -------------------------------------------------
        # CREATE PDF
        # -------------------------------------------------

        pdf_buffer = io.BytesIO()

        document = SimpleDocTemplate(

            pdf_buffer,

            pagesize=landscape(A4),

            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(

            "ReportTitle",

            parent=styles["Title"],

            alignment=TA_CENTER,

            fontSize=22,

            spaceAfter=8
        )

        subtitle_style = ParagraphStyle(

            "Subtitle",

            parent=styles["Normal"],

            alignment=TA_CENTER,

            fontSize=10,

            spaceAfter=18
        )

        heading_style = ParagraphStyle(

            "Heading",

            parent=styles["Heading2"],

            fontSize=14,

            spaceBefore=12,

            spaceAfter=8
        )

        normal_style = ParagraphStyle(

            "NormalReport",

            parent=styles["Normal"],

            fontSize=9,

            leading=13
        )

        story = []

        # -------------------------------------------------
        # TITLE
        # -------------------------------------------------

        story.append(
            Paragraph(
                "AI Business Intelligence",
                title_style
            )
        )

        story.append(
            Paragraph(
                "Sales Performance & Forecasting Report",
                subtitle_style
            )
        )

        # -------------------------------------------------
        # FILTER SUMMARY
        # -------------------------------------------------

        story.append(
            Paragraph(
                "Applied Filters",
                heading_style
            )
        )

        filter_data = [

            [
                "Category",
                category
            ],

            [
                "Region",
                region
            ],

            [
                "Month",
                month
            ]

        ]

        filter_table = Table(

            filter_data,

            colWidths=[
                45 * mm,
                65 * mm
            ]
        )

        filter_table.setStyle(
            TableStyle([

                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    6
                )

            ])
        )

        story.append(filter_table)

        # -------------------------------------------------
        # KPI SUMMARY
        # -------------------------------------------------

        story.append(
            Paragraph(
                "Key Performance Indicators",
                heading_style
            )
        )

        kpi_data = [

            [
                "Total Sales",
                "Total Orders",
                "Customers",
                "Average Order Value"
            ],

            [
                f"₹{float(kpis['total_sales']):,.2f}",

                f"{int(kpis['total_orders']):,}",

                f"{int(kpis['total_customers']):,}",

                f"₹{float(kpis['average_order_value']):,.2f}"
            ]

        ]

        kpi_table = Table(

            kpi_data,

            colWidths=[
                55 * mm,
                45 * mm,
                45 * mm,
                55 * mm
            ]
        )

        kpi_table.setStyle(
            TableStyle([

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1f2937")
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTNAME",
                    (0, 1),
                    (-1, 1),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 1),
                    (-1, 1),
                    13
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    8
                )

            ])
        )

        story.append(kpi_table)

        # -------------------------------------------------
        # CATEGORY PERFORMANCE
        # -------------------------------------------------

        story.append(
            Paragraph(
                "Category Performance",
                heading_style
            )
        )

        category_data = [
            [
                "Category",
                "Total Sales"
            ]
        ]

        for row in category_rows:

            category_data.append([

                str(row["Category"]),

                f"₹{float(row['total_sales']):,.2f}"

            ])

        category_table = Table(

            category_data,

            colWidths=[
                80 * mm,
                70 * mm
            ]
        )

        category_table.setStyle(
            TableStyle([

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "RIGHT"
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )

            ])
        )

        story.append(category_table)

        # -------------------------------------------------
        # REGION PERFORMANCE
        # -------------------------------------------------

        story.append(
            Paragraph(
                "Regional Performance",
                heading_style
            )
        )

        region_data = [
            [
                "Region",
                "Total Sales"
            ]
        ]

        for row in region_rows:

            region_data.append([

                str(row["Region"]),

                f"₹{float(row['total_sales']):,.2f}"

            ])

        region_table = Table(

            region_data,

            colWidths=[
                80 * mm,
                70 * mm
            ]
        )

        region_table.setStyle(
            TableStyle([

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "RIGHT"
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )

            ])
        )

        story.append(region_table)

        # -------------------------------------------------
        # MONTHLY PERFORMANCE
        # -------------------------------------------------

        story.append(
            Paragraph(
                "Monthly Sales Performance",
                heading_style
            )
        )

        monthly_data = [
            [
                "Month",
                "Total Sales"
            ]
        ]

        for row in monthly_rows:

            monthly_data.append([

                str(row["month"]),

                f"₹{float(row['total_sales']):,.2f}"

            ])

        monthly_table = Table(

            monthly_data,

            colWidths=[
                80 * mm,
                70 * mm
            ]
        )

        monthly_table.setStyle(
            TableStyle([

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "RIGHT"
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )

            ])
        )

        story.append(monthly_table)

        # -------------------------------------------------
        # FORECAST
        # -------------------------------------------------

        story.append(
            PageBreak()
        )

        story.append(
            Paragraph(
                "AI Sales Forecast",
                heading_style
            )
        )

        if forecast_values:

            next_forecast = float(
                forecast_values[0]
            )

            average_forecast = (
                sum(forecast_values)
                / len(forecast_values)
            )

            highest_forecast = max(
                forecast_values
            )

            forecast_summary = [

                [
                    "Next Month Forecast",
                    "Average Forecast",
                    "Highest Forecast"
                ],

                [
                    f"₹{next_forecast:,.2f}",

                    f"₹{average_forecast:,.2f}",

                    f"₹{highest_forecast:,.2f}"
                ]

            ]

            forecast_table = Table(

                forecast_summary,

                colWidths=[
                    65 * mm,
                    65 * mm,
                    65 * mm
                ]
            )

            forecast_table.setStyle(
                TableStyle([

                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey
                    ),

                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),

                    (
                        "FONTNAME",
                        (0, 1),
                        (-1, 1),
                        "Helvetica-Bold"
                    ),

                    (
                        "FONTSIZE",
                        (0, 1),
                        (-1, 1),
                        13
                    ),

                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER"
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey
                    ),

                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        8
                    )

                ])
            )

            story.append(forecast_table)

            story.append(
                Spacer(1, 10)
            )

            forecast_data = [

                [
                    "Forecast Period",
                    "Predicted Sales"
                ]

            ]

            for index, value in enumerate(
                forecast_values
            ):

                if index < len(forecast_df):

                    period = str(
                        forecast_df.iloc[
                            index,
                            0
                        ]
                    )

                else:

                    period = (
                        f"Period {index + 1}"
                    )

                forecast_data.append([

                    period,

                    f"₹{float(value):,.2f}"

                ])

            forecast_detail_table = Table(

                forecast_data,

                colWidths=[
                    80 * mm,
                    80 * mm
                ]
            )

            forecast_detail_table.setStyle(
                TableStyle([

                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey
                    ),

                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey
                    ),

                    (
                        "ALIGN",
                        (1, 1),
                        (1, -1),
                        "RIGHT"
                    ),

                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        5
                    )

                ])
            )

            story.append(
                forecast_detail_table
            )

        else:

            story.append(
                Paragraph(
                    "Forecast data is currently unavailable.",
                    normal_style
                )
            )

        # -------------------------------------------------
        # FOOTER TEXT
        # -------------------------------------------------

        story.append(
            Spacer(1, 20)
        )

        story.append(
            Paragraph(
                "Generated by AI-Powered Business "
                "Intelligence & Sales Forecasting System",
                normal_style
            )
        )

        # -------------------------------------------------
        # BUILD PDF
        # -------------------------------------------------

        document.build(story)

        pdf_buffer.seek(0)

        return send_file(

            pdf_buffer,

            as_attachment=True,

            download_name="business_sales_report.pdf",

            mimetype="application/pdf"
        )

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


# =========================================================
# RUN FLASK APPLICATION
# RAILWAY READY
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=True,
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        )
    )